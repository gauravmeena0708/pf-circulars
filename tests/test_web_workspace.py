"""Exercise migrated office workflows with real PDFs, CSVs and Word outputs."""
import io
import json
import zipfile
from types import SimpleNamespace

import fitz
from fastapi.testclient import TestClient
import pytest
from PIL import Image, ImageDraw

from web_api import app as root
from web_api import workspace


@pytest.fixture
def client():
    with workspace.LOCK:
        workspace.FILES.clear()
    return TestClient(root.app, headers={'X-Workspace-Session': 'b' * 32})


@pytest.fixture
def pdf():
    with fitz.open() as doc:
        for i in range(2):
            page = doc.new_page()
            page.insert_textbox(fitz.Rect(50, 50, 500, 500),
                f'PAGE {i+1}. ' + 'Joint declaration policy evidence for office review and member profile correction. ' * 8)
        return doc.tobytes()


def upload(client, name, data):
    response = client.post('/api/uploads', files={'file': (name, data)})
    assert response.status_code == 200, response.text
    return response.json()


def word_text(response):
    assert response.status_code == 200, response.text
    from docx import Document
    doc = Document(io.BytesIO(response.content))
    return '\n'.join([p.text for p in doc.paragraphs] +
                     [cell.text for table in doc.tables for row in table.rows for cell in row.cells])


def test_csv_profile_search_distribution_and_isolation(client):
    file = upload(client, 'claims.csv', b'office,amount\nDelhi,100\nDelhi,200\nPune,\n')
    assert file['profile']['row_count'] == 3
    assert file['profile']['null_counts']['amount'] == 1
    path = '/api/uploads/' + file['id']
    found = client.get(path + '/rows?term=Delhi&column=office').json()
    assert found['total_matches'] == 2
    assert found['distribution'] == {'Delhi': 2, 'Pune': 1}
    assert client.get(path + '/rows?column=unknown').status_code == 422
    assert client.get(path + '/download', headers={'X-Workspace-Session': 'c' * 32}).status_code == 404
    assert client.delete(path, headers={'X-Workspace-Session': 'c' * 32}).status_code == 404
    assert 'claims.csv' in word_text(client.get(path + '/download?format=docx'))
    assert client.delete(path).status_code == 200
    assert client.get(path + '/download').status_code == 404


def test_pdf_extraction_exports_and_expiry(client, pdf, monkeypatch):
    file = upload(client, 'instructions.pdf', pdf)
    assert file['page_count'] == 2
    assert '[Page 1 /' in file['text'] and '[Page 2 /' in file['text']
    path = '/api/uploads/' + file['id'] + '/download'
    assert client.get(path + '?format=original').content == pdf
    assert 'Joint declaration' in client.get(path + '?format=text').text
    assert 'Joint declaration' in word_text(client.get(path + '?format=docx'))
    assert 'instructions.pdf' in word_text(client.get(path + '?format=transcript'))
    monkeypatch.setitem(workspace.FILES[file['id']], 'created', workspace.time.monotonic()-1801)
    assert client.get(path).status_code == 404


def test_document_chat_uses_file_and_memory_only_custom_token(client, pdf, monkeypatch):
    import config
    import answer_generator
    file = upload(client, 'instructions.pdf', pdf)
    path = '/api/uploads/' + file['id']
    monkeypatch.setattr(config, 'HF_TOKEN', '')
    assert client.post(path + '/chat', json={'question': 'Summarise this file'}).status_code == 503
    calls = {}
    monkeypatch.setattr(answer_generator, 'initialize_llm', lambda hf_token: calls.update(token=hf_token) or object())
    def generate(question, chunks, model, **kwargs):
        calls.update(question=question, chunks=chunks, language=kwargs['answer_language'])
        yield SimpleNamespace(content='A supported office note.')
    monkeypatch.setattr(answer_generator, 'get_llm_answer', generate)
    response = client.post(path + '/chat', headers={'X-HF-Token': 'test-only-token'},
                           json={'question': 'Prepare a note', 'language': 'Hindi', 'mode': 'note'})
    events = [json.loads(line) for line in response.text.splitlines()]
    assert events[-1] == {'done': True}, events
    assert calls['token'] == 'test-only-token' and calls['language'] == 'Hindi'
    assert 'Joint declaration' in calls['chunks'][0]['text']
    assert 'formal office noting' in calls['question']
    text = word_text(client.get(path + '/download?format=transcript'))
    assert 'A supported office note.' in text and 'test-only-token' not in text
    assert 'test-only-token' not in repr(workspace.FILES)


def test_research_export_and_generation_are_owned_by_session(client):
    packet = client.post('/api/search', json={'question': 'Bank KYC employer approval?',
                         'policy': 'bank_kyc', 'as_of': '2025-04-03'}).json()
    body = {'packet_id': packet['packet_id'], 'answer': 'Review the source [1].'}
    assert 'Bank KYC' in word_text(client.post('/api/exports/research', json=body))
    other = {'X-Workspace-Session': 'c' * 32}
    assert client.post('/api/exports/research', headers=other, json=body).status_code == 410
    assert client.post('/api/answers/' + packet['packet_id'], headers=other, json={}).status_code == 410


@pytest.mark.parametrize('operation,count,pages', [('merge',2,4),('extract',1,1),('rotate',1,2),('compress',1,2)])
def test_pdf_operations_preserve_real_pages(client, pdf, operation, count, pages):
    response = client.post('/api/pdf-tools', data={'operation': operation, 'pages': '2'},
                            files=[('files', ('file.pdf', pdf))] * count)
    assert response.status_code == 200, response.text
    with fitz.open(stream=response.content, filetype='pdf') as doc:
        assert len(doc) == pages
        if operation == 'extract': assert 'PAGE 2' in doc[0].get_text()
        if operation == 'rotate': assert doc[0].rotation == 0 and doc[1].rotation == 90


@pytest.mark.parametrize('operation', ['split', 'render'])
def test_pdf_zip_outputs(client, pdf, operation):
    response = client.post('/api/pdf-tools', data={'operation': operation, 'dpi': 100},
                           files={'files': ('file.pdf', pdf)})
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert len(archive.namelist()) == 2
        assert all(name.endswith('.pdf' if operation == 'split' else '.png') for name in archive.namelist())


def test_images_to_pdf_and_real_searchable_ocr(client):
    image = Image.new('RGB', (1200, 500), 'white')
    from PIL import ImageFont
    font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 45)
    ImageDraw.Draw(image).text((50, 100), 'Joint declaration policy 2025', fill='black', font=font)
    source = io.BytesIO(); image.save(source, format='PNG')
    response = client.post('/api/pdf-tools', data={'operation': 'images'}, files={'files': ('scan.png', source.getvalue())})
    assert response.status_code == 200, response.text
    response = client.post('/api/pdf-tools', data={'operation': 'ocr', 'language': 'eng'},
                           files={'files': ('scan.pdf', response.content)})
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert 'extracted-document.docx' in archive.namelist()
        assert 'Joint declaration' in archive.read('extracted-text.txt').decode()
        with fitz.open(stream=archive.read('searchable.pdf'), filetype='pdf') as doc:
            assert 'Joint declaration' in doc[0].get_text()
            assert doc[0].search_for('declaration')


def test_invalid_uploads_ranges_and_request_limit(client, pdf):
    assert client.post('/api/uploads', files={'file': ('bad.exe', b'invalid')}).status_code == 422
    assert client.post('/api/uploads', files={'file': ('bad.pdf', b'invalid')}).status_code == 422
    assert client.post('/api/uploads', headers={'X-Workspace-Session': ''}, files={'file': ('a.pdf',pdf)}).status_code == 400
    assert client.post('/api/pdf-tools', data={'operation':'extract','pages':'300'}, files={'files':('a.pdf',pdf)}).status_code == 422
    response = client.post('/api/uploads', headers={'Content-Length': str(101*1024*1024)}, content=b'')
    assert response.status_code == 413
